"""Build local editable product documents and vectors from reviewed specifications.

Usage: python scripts/build_materials.py --materials DIRECTORY --runtime DIRECTORY
Requires python-docx. Complete Word pagination/render QA separately before release.
"""
import argparse
import json
import re
import sqlite3
from pathlib import Path
from xml.sax.saxutils import escape

PAGES = [
('overview','工作台概览','选择工作上下文并识别下一任务','实验选择|运行选择|任务契约数|材料版本数|实验数|观察记录数|费用上界|闭环入口','GET /api/workbench/summary; GET /api/experiments; GET /api/runs','01,09,11'),
('environment','环境检查','确认所选执行目标的依赖','执行目标|检查按钮|检查项|是否必需|就绪结果|处理动作','GET /api/doctor?target=builtin|filesystem|model','05,15'),
('catalog','MCP与任务目录','找到工具配置及可验收任务','Profile ID|固定镜像|工具白名单|Task ID|任务族|数据划分|查看任务包|编辑入口','GET /api/catalog; GET /api/tasks/{id}/package','01,02'),
('task','任务契约编辑','把任务描述转为可独立验收契约','Task ID|任务族|Profile|数据划分|公开任务说明|输入夹具JSON|输出文件|规则操作|JSON路径|预期值|协议步骤JSON|导入JSON|增加规则|预检契约|预检并登记|错误位置','POST /api/tasks/preflight; POST /api/tasks/import','01,02,06'),
('experiment','实验配置','冻结任务材料与执行预算','实验名称|任务多选|材料多选|执行模式|重复次数|运行规模|实验费用上限|创建并执行','POST /api/experiments; POST /api/experiments/{id}/batch','03,07,11'),
('comparison','实验结果比较','在可比条件下解释材料差异','来源分组|执行模式|Profile|可比状态|材料版本|接受与总数|正常任务|预期拒绝|费用上界|运行与条件','GET /api/experiments/{id}/comparison','07,11'),
('history','运行历史','检索运行并恢复工作上下文','状态筛选|运行ID|任务ID|材料版本|模式|状态|查看证据','GET /api/runs?experiment_id={id}','09,10'),
('run','运行详情与诊断','定位失败规则并查看可复验证据','运行状态|独立验收|请求与工具数|执行秒数|费用上界|工具调用|错误码|规则结果|处理建议|执行指纹|重新核验|取消|问题包下载|关联反馈','GET /api/runs/{id}; GET /api/runs/{id}/diagnosis; POST /api/runs/{id}/verification; POST /api/runs/{id}/cancel','05,06,10,11,14'),
('materials','材料与版本差异','保存有理由的不可变材料版本','父版本或Profile|名称|修改理由|接入指南|工具描述|保存新版本|版本差异','POST /api/materials; GET /api/materials/{id}/diff','02,04'),
('feedback','反馈与修订','关联问题新材料及同条件复验','当前运行|问题与证据|保存反馈|反馈ID|原反馈|新运行|决策理由|关联复验|证据状态|核验时间','POST /api/feedback; POST /api/revisions; GET /api/experiments/{id}/handoff','04,06,07,08'),
('release','发布检查与导出','检查选定版本任务集的证据','材料版本|任务集|执行模式|检查按钮|逐任务结果|总体判定|检查记录|证据导出','POST /api/releases; GET /api/releases; GET /api/releases/{id}/export','06,12'),
('developer','开发者首次任务','获得第一个可核验MCP产物','内置依赖说明|检查内置目标|执行首次任务|查看独立验收|反馈交接','GET /api/doctor?target=builtin; POST /api/experiments; POST /api/experiments/{id}/runs','04,05,09'),
('observations','使用观察与指标','保存带来源与协助程度的观察','匿名标识|来源|开始时间|结束时间|协助程度|证据引用|同意|保存观察|结果|删除观察|指标来源','POST /api/observations; GET /api/observations; DELETE /api/observations/{id}; GET /api/funnel','08,13'),
('showcase','公开案例与作品','浏览案例和可下载的脱敏证据','案例选择|六步流程|上一步|下一步|语言切换|真实报告|历史实验|精选原型|代码链接|设计入口','静态HTML/JS；GET workbench-evidence.json、report.json','16,17'),
]
SCENES = [
('接入与任务登记','维护者','接入或升级MCP工具时','手工核对工具参数、示例与运行依赖，规则分散在脚本和文档中','在登记前发现结构、路径和验收规则问题','预检定位→修正输入→登记新ID→目录查看','01,02,03,04'),
('开发者首次成功','首次接入开发者','初次安装并准备调用工具时','缺少明确起点，难以区分环境失败和任务产物错误','完成一个零模型费用的可核验任务','检查内置环境→执行task-01→查看规则→反馈交接','02,08,12'),
('材料效果实验','工具维护者','准备改变指南或工具描述时','工具可执行与模型能正确选择工具混在同一成功率中','在冻结条件下解释材料变化与任务结果','选择任务材料→冻结预算模式→执行→分组比较','05,06,07'),
('失败诊断与修订','维护者','看到失败或接到接入反馈时','错误、材料版本与复验结果缺少连续关系，陈旧通过可能误用','形成可定位、可修订、可复验的证据链','诊断→反馈→新材料→新运行→核验→关联','08,09,10'),
('发布检查与证据','维护者与案例读者','准备对外说明材料版本可用性时','报告可能遗漏任务或引用已经变化的产物','按指定材料和任务集生成可追溯发布证据','选择范围→核验最新运行→三态判定→脱敏导出','11,14'),
('使用观察与公开展示','观察者与公开读者','试用或评估作品时','演示、自动运行和真人采用容易混淆','理解案例并获得有来源的使用观察','案例浏览→真实报告→实际试用→同意及观察记录','13,14'),
]

RULES = {
'Task ID':'字符串，使用合法ID；内置ID只读，已有ID同内容重试返回原定义，修改内容使用新ID。',
'实验名称':'必填字符串，1—100字符；创建成功返回experiment_id。',
'任务多选':'1—16个不同任务；全部ID须已登记且与材料工具集匹配。',
'材料多选':'1—4个不同版本；保存的内容和哈希保持不可变。',
'重复次数':'整数1—3；规模=任务数×材料数×重复次数，最大192。',
'实验费用上限':'CNY正数，不超过全局200及既有约束；默认界面10；未知请求保留预留。',
'执行模式':'protocol按参考步骤执行；model使用指南与工具描述生成调用。',
'公开任务说明':'进入被测模型的任务上下文；包含目标及输出位置，Gold独立保存。',
'输入夹具JSON':'合法JSON对象；相对安全文件名映射到UTF-8文本；禁止路径穿越。',
'JSON路径':'JSON数组；字段校验失败定位具体规则路径；空路径表示根对象。',
'规则操作':'exists、equals、contains、set_equals、range、type；操作决定对应期望字段。',
'预期值':'按操作使用expected/type/min/max，有限数值与准确JSON类型；布尔值不冒充数值。',
'协议步骤JSON':'参考工具步骤只用于protocol；模型执行上下文隔离参考答案。',
'预检契约':'检查字段、依赖、路径、规则和匹配关系；返回valid及errors.location/code。',
'预检并登记':'先预检，再不可变登记；失败保留输入并定位错误。',
'错误位置':'语法错误提供行列与选区；规则错误链接至实际控件，aria-invalid标记。',
'名称':'材料显示名称，最多100字符；不影响材料唯一ID。',
'修改理由':'最多1000字符；关联父版本，说明要解决的具体问题。',
'接入指南':'必填1—6000字符；作为system上下文，按材料版本冻结。',
'工具描述':'按profile完整工具集合填写；每项1—1500字符，进入Function Calling定义。',
'问题与证据':'必填1—2000字符，描述步骤、预期、实际及运行引用；过滤明显凭据。',
'决策理由':'必填1—2000字符；与原反馈、新运行及材料差异对应。',
'可比状态':'按来源、模式、profile分组；未知模型或未完成为INSUFFICIENT，条件或矩阵不同为INCOMPARABLE。',
'关联复验':'同实验同任务同来源、材料改变、条件一致；原证据完整、新产物通过；保存核验指纹。',
'证据状态':'verified当前指纹一致；stale证据变化；historical缺少新式快照。',
'核验时间':'UTC ISO时间，表示关联时核验时间；历史记录保持原始结论。',
'总体判定':'存在失败为failed；缺运行、未知身份或条件不足为insufficient；全部验收通过为passed。',
'匿名标识':'1—64位字母数字下划线或连字符；用于本地关联，身份由观察证据确认。',
'来源':'observed_session、自报self_report、automation分别记录和统计。',
'开始时间':'带时区ISO时间；与结束时间共同定义该次观察区间。',
'结束时间':'带时区ISO时间且不早于开始；关联运行应已进入终态。',
'协助程度':'none、hint、guided；分别表示独立、提示和引导。',
'证据引用':'1—100字符本地安全引用，禁止凭据和路径穿越；公开导出排除。',
'同意':'保存观察要求true；撤回主体与删除单条观察分别处理。',
'保存反馈':'一次提交使用稳定幂等键；同键同内容返回原ID，同键异内容HTTP409。',
'取消':'queued可直接取消，running发出取消请求；按实际执行终止后展示结果。',
'重新核验':'重新读取产物并验证规则、manifest、任务和产物哈希；扩展验证器需显式受信任执行。',
}

RULES.update({
'实验选择':'读取GET /experiments，显示title、保存id；切换后刷新运行候选，旧运行不属于新实验时选最新运行。',
'运行选择':'仅显示所选实验下运行；标签由任务、状态和短ID组成，保存完整ID用于后续接口。',
'任务契约数':'汇总内置与已登记任务数量，按catalog当前集合计数。',
'材料版本数':'materials表版本记录数；父子材料分别计数。',
'实验数':'experiments表记录数，空实验仍计入。',
'观察记录数':'已保存观察会话条数；按来源拆分，匿名标识不直接等同真实人数。',
'费用上界':'所有charge的COALESCE(actual,reserved)之和，单位CNY；unknown请求保留预留。',
'闭环入口':'环境、实验、运行、反馈、发布导航；依赖运行的入口要求有效上下文。',
'执行目标':'builtin、filesystem、model三选一；检查按目标所需依赖组成。',
'检查按钮':'调用doctor；运行目录可写探测及本地依赖检查，不触发模型推理。',
'检查项':'显示检查id、passed、required与action；每项结果独立。',
'是否必需':'required=true的失败阻止ready；建议项提示仍保留。',
'就绪结果':'后端ready决定目标就绪；保留各失败依赖和下一处理动作。',
'处理动作':'按具体依赖显示安装、启动、登记或配置动作，用户显式处理后重新检查。',
'Profile ID':'registry中profile的不可变ID；关联镜像、工具白名单及资源限制。',
'Profile':'从已登记profile选择；任务和材料工具集合必须匹配。',
'固定镜像':'profile.image为sha256固定ID，须可在当前Docker引擎inspect；重建镜像登记新版本。',
'工具白名单':'profile.tools及实际MCP工具schema交集；未知工具在执行前拒绝。',
'任务族':'任务family用于同类场景分组，随不可变任务包保存。',
'数据划分':'exploration或holdout；沿用冻结任务划分，避免按结果回选。',
'查看任务包':'维护者读取完整任务契约，可查看Gold；模型执行使用独立的受限上下文投影。',
'编辑入口':'将所选任务包载入编辑器；修改定义使用新Task ID。',
'输出文件':'任务输出目录内安全相对文件名；对应规则file字段，路径穿越被拒绝。',
'导入JSON':'保留原始文本；语法失败定位行列；结构失败显示JSON路径；通过后载入表单。',
'增加规则':'追加一条file/op/path/expected规则；用户可移除，提交时统一预检。',
'运行规模':'即时计算选中任务数×材料数×重复次数，提交前校验非空、唯一和上限。',
'创建并执行':'先创建实验，再提交批次；创建成功后保存实验上下文，批次失败可回到该实验检查。',
'来源分组':'automation、human_unverified、human_observed分别统计；withdrawn不进入新比较组。',
'材料版本':'不可变材料ID、内容hash及父版本；显示名称仅便于识别。',
'接受与总数':'分子为succeeded且verification.passed；分母为该组已计划运行，未完成保留。',
'正常任务':'expected_error为空的任务，正常通过和计划数分别统计。',
'预期拒绝':'expected_error存在的任务，正确拒绝由独立规则决定。',
'运行与条件':'展示参与组的完整run_ids、条件哈希及balanced状态，支持追溯明细。',
'状态筛选':'按实际status筛选；终态包含succeeded、failed、cancelled、timed_out、budget_exhausted、interrupted、uncertain。',
'运行ID':'服务端生成32位ID；列表短显示，接口使用完整值。',
'任务ID':'读取运行task_id；对应冻结任务内容及provenance.task_hash。',
'模式':'读取运行mode；protocol与model分开展示。',
'状态':'读取运行status；queued/running轮询，终态停止自动轮询。',
'查看证据':'设置实验与运行上下文，进入运行详情。',
'运行状态':'同时展示执行status和独立verification，分别表示执行终态与产物验收。',
'独立验收':'verification.passed及reason/checks来源于独立验证器；无结果显示尚未完成。',
'请求与工具数':'result.requests和tools分别计数，模型请求与工具调用使用独立单位。',
'执行秒数':'result.elapsed_seconds实测墙钟耗时；未知显示破折号。',
'工具调用':'按已记录顺序展示工具名、错误及参数指纹；原始私密内容排除公开导出。',
'错误码':'读取result.error_code和diagnosis.code，绑定真实执行或规则失败。',
'规则结果':'逐条显示规则位置及passed；总体通过由完整规则集决定。',
'处理建议':'由确定性错误类别映射产生，维护者查看证据后决定修改。',
'执行指纹':'展示任务、材料、执行器、验证器、模型配置、工具schema和依赖等来源。',
'问题包下载':'白名单JSON包含目标运行及冻结配置；排除私人fixture和原始对话。',
'关联反馈':'进入反馈页并保留当前run_id，确保问题引用准确。',
'父版本或Profile':'选父材料复制内容，或选profile创建首个材料；工具集合由所选项决定。',
'保存新版本':'校验后创建新material ID和hash；保存name、reason、parent。',
'版本差异':'统一格式JSON对比父版本，显示guide和descriptions变化。',
'当前运行':'读取全局选择的完整run_id；提交反馈同时校验experiment_id。',
'反馈ID':'不可变反馈记录，关联原运行和提交时间。',
'原反馈':'从当前实验handoff.feedback选择，作为修订old_run依据。',
'新运行':'候选为当前实验运行；提交时后端核对任务、来源、条件、材料和产物。',
'任务集':'1个或多个已登记任务，无重复；逐任务检查指定材料与模式的最新对应运行。',
'逐任务结果':'passed/failed/insufficient、reason和run_id分别呈现。',
'检查记录':'不可变release_checks快照，保留created、material_hash、任务检查和比较状态。',
'证据导出':'返回adoptlab-release-v1白名单快照，Content-Disposition提示JSON下载。',
'内置依赖说明':'Python3.11及MCP依赖，stdio和CPU执行，无Docker及模型凭据要求。',
'检查内置目标':'调用doctor builtin，失败返回具体检查项与处理方式。',
'执行首次任务':'创建实验并执行task-01/B/protocol，来源human_unverified；实际观察另登记。',
'查看独立验收':'进入当前运行详情，查看规则、产物和复验入口。',
'反馈交接':'将开发者所遇问题关联当前运行，维护者获得可定位证据。',
'保存观察':'要求同意、有效时间、终态运行及证据引用；同参与者/运行/起止/来源去重。',
'结果':'accepted由所关联运行成功且独立通过计算；耗时由观察时间差计算。',
'删除观察':'删除指定观察记录，影响观察统计；主体撤回使用独立接口。',
'指标来源':'本地实例漏斗、自动化、观察者声明及自报分别说明口径。',
'案例选择':'三案例切换当前内容，演示状态与真实protocol报告独立展示。',
'六步流程':'契约→环境→诊断→材料→复验→发布；演示数据明确标记。',
'上一步':'在当前案例步骤索引范围内回退，第一步禁用。',
'下一步':'在当前案例步骤索引范围内前进，最后一步禁用。',
'语言切换':'中文英文完整切换；固定JSON字段与技术标识保持原值。',
'真实报告':'展示保存的真实协议运行及发布结果，提供版本化JSON下载。',
'历史实验':'保留144轮矩阵及旧版报告链接和原始分母。',
'精选原型':'公开脱敏SVG，页面编号对应PRD；全文Word在本地交付。',
'代码链接':'跳转Oscar-Williams/adoptlab公开仓库，查看版本与复现方法。',
'设计入口':'链接现有Figma文件；访问权限和云端保存状态单独说明。',
})


def section_map(text):
    parts=re.split(r'(?m)^## (\d+)\. ',text)
    return {int(parts[i]):parts[i+1] for i in range(1,len(parts),2)}


def md_table(text, starts):
    return [[c.strip() for c in line.strip().strip('|').split('|')] for line in text.splitlines() if line.startswith('|') and any(line.startswith('|'+x) for x in starts)]


def build_docs(out, runtime, plan):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    sections=section_map(plan)
    evidence=json.loads((runtime/'v04-validation.json').read_text())
    report=json.loads((Path(__file__).resolve().parents[1]/'public-site/report.json').read_text())

    def new(title,code,lead):
        d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.left_margin=sec.right_margin=Inches(.75)
        for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
            st=d.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
            st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
            for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
        d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(6)
        d.styles['Normal'].paragraph_format.line_spacing=1.15
        d.add_heading(title,0);d.add_paragraph(code+' · V1.0 · 产品0.5.0 · 2026-10-12 · 维护者 Oscar')
        d.add_paragraph(lead);d.add_paragraph('状态：实施与验证规格。真实观察结果待实际参与后更新；功能验证以关联证据为准。')
        d.add_paragraph('目录').runs[0].bold=True;p=d.add_paragraph();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');p._p.append(f)
        footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run(code+'  |  ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
        d.core_properties.title=title;d.core_properties.author='Oscar';d.core_properties.subject='MCP接入验证与材料修订';return d
    def table(d,headers,rows):
        t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.style='Table Grid'
        widths=([1.05,1.05,2.55,2.35] if headers[0]=='编号' else [1.55,1.35,.8,.8] if headers==['字段','类型','必填','主键'] else [7/len(headers)]*len(headers))
        if headers[0]=='ND':widths=[.85,1.65,1.5,1.5,1.5]
        for col,width in zip(t.columns,widths):col.width=Inches(width)
        for i,c in enumerate(t.rows[0].cells):c.text=headers[i]
        for row in rows:
            cells=t.add_row().cells
            for c,v in zip(cells,row):c.text=str(v).replace('`','').replace('**','')
        for ri,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            if ri==0:trpr.append(OxmlElement('w:tblHeader'))
            for i,c in enumerate(row.cells):
                c.width=Inches(widths[i]);pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8F0F2' if ri==0 else ('F7F9FA' if ri%2==0 else 'FFFFFF'));pr.append(sh)
                margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    v=OxmlElement('w:'+side);v.set(qn('w:w'),'100');v.set(qn('w:type'),'dxa');margins.append(v)
                pr.append(margins)
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.1
                    if ri==0:p.paragraph_format.keep_with_next=True
                    for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
        d.add_paragraph();return t
    def prose(d,title,text):d.add_heading(title,2);d.add_paragraph(text)
    def include(d,text):
        lines=text.splitlines();pending=[]
        def flush():
            if pending:
                valid=[r for r in pending if not all(re.fullmatch(r'[-: ]+',x or '-') for x in r)]
                if valid:table(d,valid[0],valid[1:])
                pending.clear()
        for line in lines:
            if line.startswith('|'):pending.append([x.strip() for x in line.strip('|').split('|')]);continue
            flush()
            if line.startswith('### '):d.add_heading(line[4:],2)
            elif line.strip():d.add_paragraph(line.replace('**','').replace('`',''))
        flush()
    def evidence_section(d):
        d.add_paragraph('EVID-NEW：实验 '+evidence['experiment_id']+'，三个案例各两次真实protocol运行，反馈、修订关联与发布检查完整。')
        table(d,['案例','原运行','新运行','发布结果'],[[c['task_id'],c['old_run'],c['new_run'],c['release']['verdict']] for c in evidence['cases']])
        d.add_paragraph('EVID-H144：冻结2×2指南与工具描述矩阵，12任务×4条件×3次，共144轮；AA 23/36、AB 35/36、BA 36/36、BB 36/36。结果为该离线任务集的观察差异。')
        table(d,['任务','材料','重复','真实结果'],[[x['task_id'],x['material'],x['trial'],x['reason']] for x in report['cases'] if x['task_id']=='task-09'])
        d.add_paragraph('真实观察记录基线为0。首次成功时间、协助率、采用率和满意度在取得合格观察后计算。')

    d=new('AdoptLab需求文档','D1','维护者需要确认开发者能否依据接入材料完成MCP任务，并将失败转化为可核验修订。本文定义用户问题、六组业务场景、需求来源和结果口径，作为解决方案及PRD的输入。')
    d.add_heading('1 目的与需求依据',1)
    d.add_paragraph('用户价值是减少接入过程中证据与材料的脱节：维护者能查到失败位置、对应版本和同条件复验，开发者能走通首次任务。当前需求依据为现有工程、历史离线失败与产品评审；外部需求频率与收益仍需真实研究。')
    evidence_section(d)
    d.add_heading('2 角色与业务边界',1)
    table(d,['角色','任务','交接产物'],[['维护者','登记契约、诊断、修订与检查发布','任务包、材料版本、证据报告'],['首次接入开发者','环境检查、首次任务、反馈','运行ID与实际问题'],['观察者','取得同意并记录协助和结果','匿名观察及本地证据'],['公开读者','浏览案例、下载报告、理解实现','精选原型与公开报告']])
    d.add_heading('3 六组需求场景',1)
    for i,(name,actor,when,pain,goal,flow,pages) in enumerate(SCENES,1):
        d.add_heading(f'SC-{i:02d} {name}',2)
        records=[('目的',goal),('场景',f'{actor}在{when}，通过本地工作台或公开案例完成相关任务；发生频率待真实观察。'),('需求描述',pain+'。期望结果：'+goal+'。'),('涉及人员',actor+'负责该环节；实际维护者Oscar承担产品、工程及测试职责。'),('预计上线','0.4为历史基线；0.5按功能、CI、预览和正式站点门槛发布，日期以实际部署记录为准。'),('第三方依赖','内置任务依赖Python及stdio MCP；Filesystem依赖固定Linux镜像；model依赖模型API与费用约束。'),('数据来源','已登记任务、fixture、独立Gold、材料版本、运行及反馈；使用UTF-8和JSON，内容由维护者版本化维护。观察由实际参与过程提供。'),('指标','按协议验收、模型任务效果、工作流完整性、真实观察四组统计；分母和未知项见第5节。'),('安全','本地执行采用工具白名单、路径限制和容器隔离；凭据保留本地，公开输出使用字段白名单。'),('兜底','环境失败返回处理动作；输入错误定位字段；超时及预算终止保留证据；修订由维护者决定后显式复验。')]
        for title,text in records:prose(d,title,text)
        d.add_paragraph('业务路径：'+flow+'；对应UI-'+pages.replace(',', '、UI-')+'。')
    d.add_heading('4 需求池与优先级',1)
    for i,scene in enumerate(SCENES,1):
        table(d,['字段','内容'],list(zip(['ID','模块','名称','描述或链接','来源','类型','提出时间','优先级','状态','版本','负责人','备注'],[f'REQ-{i:02d}',scene[0],scene[4],f'SC-{i:02d}及PRD对应FR','历史运行、既有工作台及需求评审','功能与证据可信性','2026-10-12','P0' if i in [1,3,4,5] else 'P1','实现与验收','0.5','Oscar','真实采用待观察；不适用扩展按方案保留理由'])))
    d.add_heading('5 指标与数据要求',1);include(d,sections[10].split('### 10.2')[1].split('### 10.3')[0])
    d.add_heading('6 研究与确认方式',1)
    d.add_paragraph('围绕最近一次接入询问所用工具、具体失败、排查过程、材料、帮助来源及已有产物。观察任务固定为内置首次任务，再按参与者经验选择文档配置案例。记录开始结束、协助、任务ID、运行ID、结果、退出和同意。参与者可以中止，原始身份关联仅本地保留。')
    d.add_paragraph('需求关闭：六场景能够定位到功能、验收和证据；工程指标由记录验证，用户收益由真实观察验证。')
    d.save(out/'AdoptLab需求文档.docx')

    d=new('AdoptLab需求解决方案','D2','选用本地受约束执行与线上证据展示，围绕任务契约、材料版本、失败证据和复验建立连续工作流。本文说明方案取舍、技术可行性、POC及真实案例结果。')
    for title,body in [('1 问题与目标','目标是让维护者把接入失败转化为有证据的材料修改，并确认修改后的任务验收。六组需求及研究状态见D1。'),('2 旧流程与目标流程','旧流程：编辑文档→手工运行脚本→分散保存日志→口头判断修复。目标流程：预检契约→冻结材料与条件→执行→独立规则验收→反馈→不可变修订→同条件复验→发布证据。')]:d.add_heading(title,1);d.add_paragraph(body)
    d.add_heading('3 候选方案比较',1)
    table(d,['方案','能力依据','适合条件','选择'],[['手工脚本','可执行工具与自定义断言，关联关系依赖维护者','少量一次性检查','保留脚本作为可复现入口'],['Promptfoo','官方MCP集成支持工具评测与结果查询','已有评测配置和数据集','作为比较对象，不重复实现通用评测平台'],['Langfuse','官方数据集、实验与追踪','已有观测基础设施','可选观测集成，避免成为本地运行硬依赖'],['AdoptLab','任务契约、材料版本、失败证据和修订关系连续保存','MCP维护者接入材料迭代','深化现有工程，独立规则验收和本地运行']])
    d.add_heading('4 技术与产品取舍',1);include(d,sections[6].split('### 6.2')[1].split('### 6.3')[0])
    d.add_heading('5 需求架构与节点',1)
    d.add_paragraph('维护者的契约登记事件产生任务版本；执行事件产生运行与产物；反馈事件关联原运行；修订事件关联新材料与新运行；发布事件冻结任务集检查。开发者使用同一任务与验收链路，公开读者访问脱敏投影。')
    include(d,sections[9].split('### 9.2')[1].split('### 9.3')[0])
    d.add_heading('6 技术预研与POC',1)
    steps=['限定三案例及验收目标','准备fixture与独立Gold','识别模型身份和产物完整性风险','冻结MCP工具schema及白名单','固定Python、容器和模型配置','判断检索适用性：指定小文件无需向量检索','保留protocol与历史model基线','区分人工判断、自动化与真实观察','回溯真实Bad Case','控制材料变量并独立复验','测量耗时、请求、工具与费用','形成继续、调整或暂停决定']
    for i,step in enumerate(steps,1):prose(d,f'POC-{i:02d} {step}','输入由对应任务、材料和运行指纹固定。执行结果关联EVID，异常保留原始状态。'+('历史task-09存在output_mismatch；材料矩阵结果保留逐次记录，不以协议成功推断模型收益。' if i==9 else '验收采用PRD对应AC及独立规则。'))
    d.add_heading('7 三案例验证',1);evidence_section(d)
    d.add_heading('8 风险与决定',1);include(d,sections[11].split('\n',1)[1])
    d.add_heading('9 成本性能及继续条件',1)
    d.add_paragraph('protocol不产生模型费用。model每次最多8个请求、单次响应最多8个工具调用、单次输出1024 token；输入按UTF-8字节预算控制。未知usage保留预留费用。继续条件为独立验收可执行、费用在原约束内、证据支持结论。真实采用效益待观察。')
    d.add_heading('10 直接依据',1);d.add_paragraph('课程1—8、PRD_Ready、现有任务和独立验证器、report.json历史矩阵、新版三案例记录；Promptfoo https://www.promptfoo.dev/docs/integrations/mcp-server/；Langfuse https://langfuse.com/docs/evaluation/experiments/datasets。')
    d.save(out/'AdoptLab需求解决方案.docx')

    d=new('AdoptLab产品需求文档PRD','D3','本规格服务研发、测试和维护，定义14类界面、AI执行边界、数据关系和验收。源自D1用户问题与D2方案决策；元素与节点编号与两份SVG对应。')
    d.add_heading('1 文档控制与评审',1);table(d,['版本','修改内容','时间','责任'],[['0.4基线','原工作台与三案例协议闭环','2026-10-07','Oscar'],['0.5规格','分组比较、实时核验、幂等、字段定位及迁移','2026-10-12','Oscar']]);d.add_paragraph('评审按产品、工程、测试、数据和设计视角执行，由实际维护者承担；用户确认实施范围。功能完成状态以证据和测试结果回填。')
    d.add_heading('2 需求概述',1);table(d,['需求','场景','目标'],[[f'REQ-{i:02d}',x[0],x[4]] for i,x in enumerate(SCENES,1)])
    d.add_heading('3 产品定义与全局规则',1)
    d.add_paragraph('本地维护者拥有任务登记、执行与材料管理权限；公开站点提供静态案例和报告。实验与运行选择保存在当前浏览器；任务ID与材料内容不可变。全局状态区分加载、空数据、输入错误、依赖缺失、运行中、失败、取消、超时、预算不足、不可比和复验通过。')
    d.add_paragraph('主路径：维护者UI-02→04→05→08→09→10→11；开发者UI-12→08→10；公开读者UI-14。错误返回输入或依赖所在节点，运行重试由明确操作触发。')
    d.add_heading('4 功能详细规格',1)
    for i,(key,name,goal,controls,api,ac) in enumerate(PAGES,1):
        d.add_heading(f'FR-{i:02d} {name}',2);d.add_paragraph('角色：'+('公开读者' if i==14 else '开发者' if i==12 else '维护者')+'。目标：'+goal+'。触发：打开对应导航或上游任务入口。前置：本地服务可用，依赖实验/运行的操作已选择上下文。')
        d.add_paragraph('主流程：读取关联对象→显示当前状态→提交明确操作→后端校验→显示结果及下一入口。输入错误保留内容并定位；运行异常保留证据和费用状态；读取失败提供恢复入口。后置：读取保持记录，写入返回持久化对象ID。')
        rows=[]
        for j,label in enumerate(controls.split('|'),1):
            rule=RULES.get(label,'从对应API读取实际字段与版本；空值显示明确状态，长ID可复制或换行。')
            actions={'检查按钮','预检契约','预检并登记','创建并执行','保存新版本','保存反馈','关联复验','检查内置目标','执行首次任务','保存观察','删除观察','重新核验','取消'}
            editable={'Task ID','任务族','Profile','数据划分','公开任务说明','输入夹具JSON','输出文件','规则操作','JSON路径','预期值','协议步骤JSON','导入JSON','实验名称','任务多选','材料多选','执行模式','重复次数','实验费用上限','父版本或Profile','名称','修改理由','接入指南','工具描述','问题与证据','原反馈','新运行','决策理由','材料版本','任务集','匿名标识','来源','开始时间','结束时间','协助程度','证据引用','同意'}
            navigable={'闭环入口','处理动作','查看任务包','编辑入口','查看证据','关联反馈','查看独立验收','反馈交接','六步流程','上一步','下一步','语言切换','代码链接','设计入口','案例选择'}
            interaction=('点击后禁用重复提交；成功呈现返回对象或操作结果，失败保留输入与上下文。' if label in actions or (key=='release' and label=='检查按钮') else
                         '获得焦点后编辑或选择；提交前按规则校验，错误定位此控件。'+('草稿在当前浏览器恢复。' if key=='task' else '') if label in editable and key not in ['catalog','comparison','history'] else
                         '点击进入目标页面或状态；保留所属案例、实验或运行上下文。' if label in navigable else
                         '通过下载链接取得版本化文件；失败显示原因并允许重新请求。' if label in {'问题包下载','证据导出','真实报告','历史实验','精选原型'} else
                         '选择后刷新对应运行或实验上下文，并保存当前选择。' if label in {'实验选择','运行选择','状态筛选','执行目标'} else
                         '只读展示后端返回值；无结果显示明确空态，长标识支持换行。')
            rows.append([f'UI-{i:02d}-E{j:02d}',label,rule,interaction])
        table(d,['编号','字段名称','规则解释','交互逻辑'],rows)
        d.add_paragraph('接口：'+api+'。验收：AC-'+ac.replace(',', '、AC-')+'。页面：UI-'+f'{i:02d}'+'。')
    d.add_heading('5 AI能力与节点规格',1);include(d,sections[10].split('### 10.1')[1].split('### 10.2')[0]);include(d,sections[9].split('### 9.2')[1].split('### 9.3')[0])
    d.add_heading('6 非功能需求',1)
    for title,text in [('性能','测量三案例总耗时、请求和工具数；默认episode180秒，HTTP45秒，MCP读20秒；目标以实际任务规模验收。'),('可靠性','单运行认领、事务幂等、取消与中断保留状态；数据库迁移前一致性备份；恢复时不自动重放未知费用请求。'),('安全与隐私','loopback Host与Origin校验、JSON请求体20000字节上限、工具白名单、容器网络隔离和只读根；公开字段白名单。'),('兼容与可访问','Windows本机、Linux CI及Filesystem分别验证；中文英文、键盘焦点、错误提示和1440/1024/390布局检查。'),('维护与可观察','状态、错误码、运行指纹、费用预留及发布记录可查；保留历史报告、迁移备份和依赖版本。')]:prose(d,title,text)
    d.add_heading('7 数据接口与依赖',1)
    with sqlite3.connect(runtime/'adoptlab.db') as c:
        for (name,) in c.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"):
            d.add_heading(name,2);table(d,['字段','类型','必填','主键'],[[r[1],r[2],bool(r[3]),bool(r[5])] for r in c.execute('PRAGMA table_info('+name+')')])
    d.add_paragraph('schema 3增量迁移到4：新增submissions和revision_evidence。submissions以scope/key唯一约束，保存载荷哈希和响应ID；revision_evidence关联修订并随其撤回删除。其他实体和历史报告可读。')
    d.add_paragraph('POST /api/feedback及/revisions接受可选idempotency_key，8—64位安全字符；同键同载荷返回原ID，冲突409。修订读取附verification状态与verified_at。comparison新增comparison_sets，旧字段保留兼容。')
    from urllib.request import build_opener, ProxyHandler
    api=json.load(build_opener(ProxyHandler({})).open('http://127.0.0.1:8780/openapi.json', timeout=20))
    d.add_heading('请求字段与校验',2)
    for model,definition in api['components']['schemas'].items():
        if model in ['HTTPValidationError','ValidationError']:continue
        d.add_heading(model,3)
        rows=[]
        for field,rule in definition.get('properties',{}).items():
            constraints={k:v for k,v in rule.items() if k not in ['title','description']}
            rows.append([field,'是' if field in definition.get('required',[]) else '否',json.dumps(constraints,ensure_ascii=False)])
        table(d,['字段','必填','类型 默认值与校验'],rows)
    d.add_heading('接口行为与副作用',2)
    for route,methods in api['paths'].items():
        if not route.startswith('/api/'):continue
        for method,operation in methods.items():
            if method not in ['get','post','delete','put']:continue
            d.add_paragraph(method.upper()+' '+route+' · '+operation.get('summary','')+' · 响应状态 '+', '.join(operation.get('responses',{})))
    d.add_heading('8 埋点指标与实验',1);include(d,sections[10].split('### 10.2')[1].split('### 10.4')[0])
    d.add_heading('9 测试验收与发布',1)
    for row in md_table(sections[14],['AC-']):
        d.add_heading(row[0],2);d.add_paragraph('Given：准备'+row[1]+'对应输入、运行和依赖。When：执行相应操作并保留响应。Then：'+row[2]+'。')
    evidence_section(d)
    d.add_paragraph('发布顺序：本地测试→功能分支→Windows/Linux CI与容器检查→站点预览→合并→正式域名及下载校验。恢复使用上一已验证代码与一致性数据备份，历史报告保留。')
    d.add_heading('10 运营与持续维护',1)
    d.add_paragraph('维护者以失败任务和反馈为入口，每次变更记录理由、父版本和复验。开发者培训从内置首次任务开始；观察者记录真实协助与结果。新工具、模型或验证器版本变化触发新的条件组。发布后检查网站可访问、下载正确、错误与预算状态。')
    d.add_heading('11 追踪与开放事项',1)
    table(d,['需求','功能与页面','验收范围'],[[f'REQ-{i:02d}',x[6], '对应FR列明的AC'] for i,x in enumerate(SCENES,1)])
    d.add_paragraph('真实采用待取得观察；模型改进使用历史冻结矩阵，新增费用需遵守现有预算；Figma云端保存状态单独确认。原型展示目标和当前状态均通过明确标签区分。')
    d.add_heading('12 直接来源',1);include(d,sections[4].split('\n',1)[1]);d.add_heading('课程要求追踪',2);include(d,sections[3].split('\n',1)[1])
    d.save(out/'AdoptLab产品需求文档PRD.docx')


def svg_text(x,y,value,size=18,width=48):
    lines=[]
    for paragraph in str(value).split('\n'):
        lines.extend([paragraph[i:i+width] for i in range(0,len(paragraph),width)] or [''])
    return '<text x="%s" y="%s" font-size="%s">'%(x,y,size)+''.join('<tspan x="%s" dy="%s">%s</tspan>'%(x,0 if i==0 else size*1.5,escape(line)) for i,line in enumerate(lines))+'</text>'


def svg_root(w,h,title,body):
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}" role="img"><title>{title}</title><desc>AdoptLab 0.5 editable product specification. Text and shapes remain editable.</desc><defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="9" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 Z" fill="#64748b"/></marker></defs><rect width="100%" height="100%" fill="#f3f8f9"/><g font-family="Microsoft YaHei,Arial,sans-serif" fill="#19353f">{body}</g></svg>'


def build_vectors(out,plan):
    b=svg_text(48,64,'AdoptLab 产品设计原型 · 0.5',32)+svg_text(48,112,'维护者：02→04→05→08→09→10→11    开发者：12→08→10    公开：14',20,100)
    for i,(key,name,goal,controls,api,ac) in enumerate(PAGES,1):
        x=48+(i-1)%3*1520;y=180+(i-1)//3*1080
        b+=f'<g id="UI-{i:02d}" transform="translate({x},{y})"><rect width="1440" height="1000" rx="12" fill="white" stroke="#c9dce2"/><rect width="232" height="1000" fill="#e8f0f2"/>'
        b+=svg_text(24,44,'AdoptLab',24)+svg_text(264,44,f'UI-{i:02d} / {name}',30)+svg_text(264,92,goal,20)
        for j,p in enumerate(PAGES):b+=svg_text(24,100+j*43,f'{j+1:02d} {p[1]}',16,14)
        for j,label in enumerate(controls.split('|')):
            cx=264+(j%2)*566;cy=136+(j//2)*88
            b+=f'<g id="UI-{i:02d}-E{j+1:02d}"><rect x="{cx}" y="{cy}" width="542" height="72" rx="6" fill="#f3f8f9" stroke="#c9dce2"/>'+svg_text(cx+16,cy+26,f'E{j+1:02d}  {label}',18,30)+svg_text(cx+16,cy+53,(RULES[label].split('；')[0][:34]+'…' if len(RULES[label].split('；')[0])>34 else RULES[label].split('；')[0]),14,34)+'</g>'
        b+=svg_text(264,930,'AC-'+ac.replace(',', ' / AC-')+' · 元素规则见PRD四列表',17,64)+'</g>'
    y=180+5*1080;b+=svg_text(48,y,'关键状态与局部变化',30)
    states=[('加载','请求中→读取成功','保持上下文，显示读取状态'),('空数据','无记录→创建或首次任务','入口明确，计数按实际显示'),('输入错误','无效JSON→行列定位','保留原文，聚焦错误位置'),('依赖缺失','检查失败→处理动作','修复依赖后显式再检查'),('运行中','queued→running','轮询状态，允许取消'),('失败','规则未通过→查看证据','定位规则，提交反馈'),('取消','请求取消→cancelled','终态与未知费用分别显示'),('超时','达到上限→timed_out','保存已完成轨迹，显式复跑'),('预算不足','预留失败→budget_exhausted','查看已结算及保留费用'),('不可比','条件漂移→INCOMPARABLE','保留明细，固定条件再执行'),('复验通过','完整证据→verified','记录时间指纹，变化显示stale')]
    for k,(name,change,action) in enumerate(states):
        x=48+k%3*1520;cy=y+40+k//3*210;b+=f'<g id="STATE-{k+1:02d}"><rect x="{x}" y="{cy}" width="1440" height="180" rx="8" fill="white" stroke="#c9dce2"/>'+svg_text(x+24,cy+40,name,24)+svg_text(x+24,cy+84,change,22)+svg_text(x+24,cy+126,action,18)+'</g>'
    y+=950
    for k,width in enumerate([1024,390]):
        for j,index in enumerate([12,14]):
            x=48+j*1520+k*600;cy=y+k*850
            b+=f'<g id="UI-{index}-{width}"><rect x="{x}" y="{cy}" width="{width}" height="760" fill="white" stroke="#c9dce2"/>'+svg_text(x+20,cy+45,f'UI-{index} · {width}px',22)+svg_text(x+20,cy+95,PAGES[index-1][1],24,15)
            for n,label in enumerate(PAGES[index-1][3].split('|')[:6]):b+=f'<rect x="{x+20}" y="{cy+130+n*86}" width="{width-40}" height="64" rx="6" fill="#e8f0f2"/>'+svg_text(x+32,cy+168+n*86,label,18,20)
            b+='</g>'
    (out/'AdoptLab产品设计原型图.svg').write_text(svg_root(4608,y+1750,'AdoptLab产品设计原型图',b),encoding='utf-8')
    sections=section_map(plan);nodes=md_table(sections[9],['ND-'])

    b=svg_text(48,60,'AdoptLab 节点流程与故障处理 · 0.5',30)+svg_text(48,108,'普通节点使用矩形；判断使用菱形；AI节点采用浅青填充；人工修订由维护者决定。',20,100)
    positions={}
    def box(id,label,x,y,w=700,h=90,kind='process'):
        nonlocal b
        positions[id]=(x,y,w,h)
        b+=f'<g id="{id}">'
        if kind=='decision':
            b+=f'<polygon points="{x+w/2},{y} {x+w},{y+h/2} {x+w/2},{y+h} {x},{y+h/2}" fill="#e8f0f2" stroke="#64748b"/>'
        else:b+=f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{40 if kind=="end" else 6}" fill="{ "#dff3f0" if kind=="ai" else "white"}" stroke="#64748b"/>'
        b+=svg_text(x+w*.22 if kind=='decision' else x+30,y+36,label,16 if kind=='decision' else 18,max(12,int((w*.56 if kind=='decision' else w-60)/18)))+'</g>'
    def edge(a,z,label='',side=False):
        nonlocal b
        x,y,w,h=positions[a];xx,yy,ww,hh=positions[z]
        if not side:
            start=x+w/2;end=xx+ww/2;mid=(y+h+yy)/2
            b+=f'<path d="M{start} {y+h} V{mid} H{end} V{yy-5}" fill="none" stroke="#64748b" marker-end="url(#arrow)"/>'
            if label:b+=svg_text((start+end)/2+12,mid-8,label,16,20)
        else:
            lane=min(x,xx)-90
            b+=f'<path d="M{x} {y+h/2} H{lane} V{yy+hh/2} H{xx-5}" fill="none" stroke="#b45309" marker-end="url(#arrow)"/>'
            if label:b+=svg_text(lane+8,(y+yy)/2,label,16,15)
    box('START','开始 · 维护者准备接入验证',1200,170,kind='end')
    initial=[('ND-01','需求与任务场景'),('ND-02','环境检查与依赖'),('ND-03','契约预检与不可变登记'),('ND-04','冻结任务及Gold'),('ND-05','选择材料与工具描述'),('ND-06','冻结模型及执行配置'),('ND-07','运行入队与认领'),('ND-08','protocol / model？')]
    previous='START'
    for i,(id,label) in enumerate(initial):
        box(id,id+' '+label,1200,320+i*150,kind='decision' if id=='ND-08' else 'process');edge(previous,id);previous=id
    box('ENV-STOP','依赖或输入失败 → 修正后重新检查',120,650,700,110,kind='end');edge('ND-02','ENV-STOP','失败')
    box('ND-P1','ND-P1 protocol：参考步骤生成调用\n普通确定性处理，模型请求为0',200,1550,750,110)
    box('ND-M1','ND-M1 组装Prompt与工具schema\n指南、公开指令、允许工具；隔离Gold',2300,1550,850,110)
    edge('ND-08','ND-P1','protocol');edge('ND-08','ND-M1','model')
    box('ND-M3','ND-M3 预算与执行限制满足？',2300,1730,850,110,kind='decision');edge('ND-M1','ND-M3')
    box('STOP','终止 · 取消 / 超时 / 预算耗尽\n保存实际状态、轨迹及未知费用',3300,1800,650,110,kind='end');edge('ND-M3','STOP','不满足')
    box('ND-M2','ND-M2 AI：理解任务、选择工具与生成参数\n受约束LLM + Function Calling',2300,1950,850,110,kind='ai');edge('ND-M3','ND-M2','满足')
    box('ND-09','ND-09 参数、白名单与MCP工具执行\n普通处理：限制路径、结果大小及资源',1200,2160,700,110);edge('ND-P1','ND-09');edge('ND-M2','ND-09','tool calls')
    box('ND-M4','ND-M4 AI：将工具结果用于下一步\n提取字段、组织产物或结束任务',2300,2330,850,110,kind='ai');edge('ND-09','ND-M4','model继续')
    edge('ND-M4','ND-M3','仍有调用且未达上限',True)
    box('ND-10','ND-10 保存产物、manifest与来源指纹',1200,2510,700,110);edge('ND-09','ND-10','protocol完成');edge('ND-M4','ND-10','model完成')
    tail=[('ND-11','独立验收通过？'),('ND-12','关联错误、轨迹与处理建议'),('ND-13','提交反馈并关联原运行'),('ND-14','人工修改材料并保存不可变版本'),('ND-15','同条件显式复验；核对实际模型身份'),('ND-16','证据完整且新产物通过？'),('ND-17','发布检查通过？'),('ND-18','白名单导出证据'),('ND-19','真实观察、协助程度及反馈'),('ND-20','公开案例浏览与版本化报告')]
    previous='ND-10'
    for i,(id,label) in enumerate(tail):
        box(id,id+' '+label,1200,2690+i*170,700,110,kind='decision' if id in ['ND-11','ND-16','ND-17'] else 'process');edge(previous,id,'需要反馈' if previous=='ND-11' else '证据完整且新产物通过' if previous=='ND-16' else '通过' if previous=='ND-17' else '');previous=id
    edge('ND-11','ND-17','通过且直接检查发布',True)
    box('RELEASE-HOLD','失败或证据不足 · 保留检查结果\n补充运行或修订后显式重检',2300,4050,850,110,kind='end');edge('ND-17','RELEASE-HOLD','失败 / 不足')
    box('REWORK','证据不足或产物变化\n定位问题后显式修订与复验',220,3400,750,120,kind='end');edge('ND-16','REWORK','失败 / 不足');edge('REWORK','ND-14','处理后',True)
    box('END','结束 · 发布证据与持续维护',1200,4470,700,90,kind='end');edge('ND-20','END')
    box('GOLD','独立Gold与验证规则\n输入字段、集合、范围及引用位置',200,2750,750,100);edge('GOLD','ND-11','仅验收使用')
    for i,(name,detail) in enumerate([('C1 内置记录','输入记录 → MCP读取/转换 → 写产物 → 字段与集合验收'),('C2 Filesystem证据','固定镜像 → read_text_file → write_file → 证据引用规则'),('C3 配置与引用','配置文档 → endpoint、retry_limit、citation_line → 逐字段验收')]):
        y=3000+i*250
        box('CASE-'+str(i+1),name+'\n'+detail+'\n均连接ND-13至ND-18的修订与发布闭环。',2300,y,1500,160)
    base=4740
    for i,row in enumerate(nodes):
        x=48+i%2*2000;y=base+i//2*650
        b+=f'<g id="CARD-{escape(row[0])}"><rect x="{x}" y="{y}" width="1940" height="610" rx="8" fill="white" stroke="#c9dce2"/>'
        labels=['节点','类别与类型','输入与输出','问题与证据','优化与处置']
        offset=42
        for label,cell in zip(labels,row):
            value=label+'：'+cell
            b+=svg_text(x+24,y+offset,value,20,85)
            offset+=max(1,(len(value)+84)//85)*30+22
        b+=svg_text(x+24,y+offset,'验收：按对应AC核验输入、输出与分支；人工介入由维护者显式决定。',18,95)+'</g>'
    (out/'AdoptLab节点流程图.svg').write_text(svg_root(4050,base+((len(nodes)+1)//2)*650+60,'AdoptLab节点流程图',b),encoding='utf-8')



def main():
    p=argparse.ArgumentParser();p.add_argument('--materials',type=Path,required=True);p.add_argument('--runtime',type=Path,required=True);a=p.parse_args()
    plan=(a.materials/'AdoptLab材料执行方案.md').read_text(encoding='utf-8')
    build_docs(a.materials,a.runtime,plan);build_vectors(a.materials,plan)
    print('Built three editable Word documents and two SVG specifications.')


if __name__=='__main__':main()
